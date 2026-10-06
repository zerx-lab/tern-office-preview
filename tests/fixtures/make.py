#!/usr/bin/env python3
"""Writes the sample documents in this folder (committed; rerun only to change them).

    uv run --with python-docx --with openpyxl --with pillow python tests/fixtures/make.py

sample.docx: headings, styled runs (CJK too), numbered and bulleted lists, a table with merged
and shaded cells, a picture, a page break, header and footer.
sample.xlsx: number/date/percent/currency formats, fonts, fills, borders, merges, wrap, column
widths, frozen header row, a hidden sheet, a long sheet split into parts, a picture.
"""

import datetime
import io
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw

HERE = pathlib.Path(__file__).resolve().parent


def picture() -> io.BytesIO:
    img = Image.new("RGB", (240, 120), (68, 114, 196))
    d = ImageDraw.Draw(img)
    d.rectangle((20, 20, 220, 100), outline=(255, 255, 255), width=4)
    d.ellipse((90, 35, 150, 85), fill=(255, 192, 0))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def shade(cell, hex_fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def make_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "Office 预览 · 页眉"
    sec.footer.paragraphs[0].text = "页脚文字"
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Office 预览示例文档", 0)
    doc.add_heading("1. 文字样式", 1)
    p = doc.add_paragraph("普通文字，")
    p.add_run("粗体").bold = True
    p.add_run("、")
    p.add_run("斜体").italic = True
    p.add_run("、")
    p.add_run("下划线").underline = True
    p.add_run("、")
    r = p.add_run("红色 14pt")
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.size = Pt(14)
    p.add_run("、")
    r = p.add_run("高亮")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run("、")
    r = p.add_run("删除线")
    r.font.strike = True
    p.add_run("，以及 H")
    r = p.add_run("2")
    r.font.subscript = True
    p.add_run("O 与 E=mc")
    r = p.add_run("2")
    r.font.superscript = True
    p.add_run("。")

    c = doc.add_paragraph("居中的段落，首行没有缩进。")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ind = doc.add_paragraph(
        "这一段首行缩进两个字符，用来检查 first-line indent 的效果。较长的文字会自动换行，"
        "换行后的第二行应当回到左边距，而不是继续缩进。The quick brown fox jumps over the lazy dog."
    )
    ind.paragraph_format.first_line_indent = Pt(21)
    ind.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("2. 列表", 1)
    for t in ["第一项", "第二项", "第三项"]:
        doc.add_paragraph(t, style="List Number")
    for t in ["苹果", "香蕉"]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph("二级项目", style="List Bullet 2")

    doc.add_heading("3. 表格", 1)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads = ["项目", "数量", "备注"]
    for i, h in enumerate(heads):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        shade(cell, "D9E1F2")
    rows = [("设备 A", "12", "已到货"), ("设备 B", "3", ""), ("合计", "15", "")]
    for r_i, row in enumerate(rows, start=1):
        for c_i, v in enumerate(row):
            table.cell(r_i, c_i).text = v
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(3, 0).merge(table.cell(3, 1)).text = "合计 15"

    doc.add_heading("4. 图片", 1)
    doc.add_picture(picture(), width=Inches(2.5))
    doc.add_paragraph("图 1 示例图片").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("5. 分页之后", 1)
    for i in range(1, 41):
        doc.add_paragraph(f"第 {i} 段：用于让文档足够长，以便检查分部分显示。" * 2)

    doc.core_properties.title = "Office 预览示例"
    doc.save(HERE / "sample.docx")


def make_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    thin = Side(style="thin", color="000000")
    head_fill = PatternFill("solid", fgColor="4472C4")
    heads = ["日期", "名称", "金额", "占比", "数量", "备注"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    base = datetime.date(2024, 1, 1)
    for r in range(2, 1202):
        ws.cell(row=r, column=1, value=base + datetime.timedelta(days=r - 2)).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=f"项目 {r - 1}")
        ws.cell(row=r, column=3, value=(r * 1234.5) % 99999 - 20000).number_format = '"¥"#,##0.00;[Red]-"¥"#,##0.00'
        ws.cell(row=r, column=4, value=(r % 100) / 100).number_format = "0.0%"
        ws.cell(row=r, column=5, value=r * 3)
    ws["F2"] = "这是一段很长的备注文字，应该在单元格内自动换行显示，而不是被截断。"
    ws["F2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B3"].font = Font(italic=True, underline="single", color="C00000")
    ws["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["E5"] = 1234567.891
    ws["E5"].number_format = "#,##0.00"
    ws["E6"] = 0.000012345
    ws["E6"].number_format = "0.00E+00"
    ws["E7"] = True
    ws["E8"] = datetime.datetime(2024, 3, 5, 14, 30, 15)
    ws["E8"].number_format = "yyyy/m/d h:mm:ss"
    ws.merge_cells("G2:H3")
    ws["G2"] = "合并单元格"
    ws["G2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["G2"].fill = PatternFill("solid", fgColor="E2EFDA")
    ws["G5"] = "超长文本会溢出到右侧的空单元格里，就像 Excel 一样"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["F"].width = 24
    ws.freeze_panes = "A2"

    s2 = wb.create_sheet("说明")
    s2["A1"] = "第二个工作表"
    s2["A1"].font = Font(size=16, bold=True)
    s2["A3"] = "图片："
    img = XlImage(picture())
    s2.add_image(img, "B4")

    s3 = wb.create_sheet("隐藏表")
    s3["A1"] = "隐藏的工作表"
    s3.sheet_state = "hidden"
    wb.save(HERE / "sample.xlsx")


if __name__ == "__main__":
    make_docx()
    make_xlsx()
    print("wrote", HERE / "sample.docx", HERE / "sample.xlsx")
